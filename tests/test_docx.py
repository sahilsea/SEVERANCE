"""Unit tests for deliver/docx.py.

Verifies:
1. Three-way classification stamping on answered response.
2. Binary Word (.docx) document integrity.
3. Filename generation stamped with classification tier.
4. Refusal to generate deliverable report for an abstained response.
"""

import io
import pytest
from docx import Document
from contracts import (
    AskResponse,
    Citation,
    Compartment,
    Label,
    Tier,
)
from deliver.docx import build_report, get_report_filename


def test_build_report_answered():
    """Build Word report for answered query with 3-way redundant classification stamping."""
    response = AskResponse(
        status="answered",
        answer="The emergency depressuring system is actuated by pushing PB-201 in the control room.",
        citations=[
            Citation(
                doc_id="sop-101",
                page=4,
                quote="emergency depressuring system is actuated by pushing PB-201 in the control room",
            )
        ],
        denials=[],
        effective_label=Label(
            tier=Tier.SECRET,
            compartments=frozenset([Compartment.TECHNICAL]),
        ),
        ledger_row_id=42,
    )

    filename = get_report_filename(response)
    assert filename == "SEVERANCE_Report_Row42_SECRET.docx"

    docx_bytes = build_report(
        response=response,
        question="How is emergency depressuring initiated?",
    )
    assert len(docx_bytes) > 1000

    # Re-open docx to verify structure and stamping
    doc = Document(io.BytesIO(docx_bytes))

    # 1. Header stamp verified
    section = doc.sections[0]
    header_text = section.header.paragraphs[0].text
    assert "CLASSIFICATION: SECRET" in header_text
    assert "TECHNICAL" in header_text

    # 2. Footer stamp verified
    footer_text = section.footer.paragraphs[0].text
    assert "CLASSIFICATION: SECRET" in footer_text

    # 3. Top banner verified
    banner_text = doc.paragraphs[0].text
    assert "CLASSIFICATION: SECRET" in banner_text


def test_build_report_abstained_raises_error():
    """Attempting to generate a deliverable report for an abstained query must fail."""
    response = AskResponse(
        status="abstained",
        answer="Access denied under two-axis security gate.",
        citations=[],
        denials=[],
        effective_label=Label(tier=Tier.PUBLIC),
        ledger_row_id=43,
    )

    with pytest.raises(ValueError, match="Cannot generate deliverable report for an abstained"):
        build_report(
            response=response,
            question="What are the secret formulas?",
        )
