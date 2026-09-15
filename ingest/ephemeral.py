"""Ephemeral, session-scoped upload handling for ad-hoc user files.

NON-NEGOTIABLE DESIGN PRINCIPLE:
Uploaded files are the user's OWN content, not governed corpus documents.
They are NEVER written to corpus/manifest.json, NEVER assigned a compartment
or tier, and NEVER subject to the two-axis security gate in trust/labels.py --
they exist only in memory, scoped to the uploading person_id, for the
lifetime of this server process. Answering questions about them is closer to
a personal document assistant than to the governed RTI corpus workbench.
"""

from __future__ import annotations

import base64
import io
import uuid
from dataclasses import dataclass, field
from typing import Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "webp", "bmp"}


@dataclass
class EphemeralImage:
    """One image extracted from an upload, base64-encoded for the vision model."""
    label: str
    base64_data: str
    mime_type: str


@dataclass
class EphemeralTextChunk:
    """One page/slide-worth of text extracted from an upload."""
    label: str
    text: str


@dataclass
class EphemeralUpload:
    upload_id: str
    filename: str
    person_id: str
    text_chunks: list[EphemeralTextChunk] = field(default_factory=list)
    images: list[EphemeralImage] = field(default_factory=list)


# In-memory only: never persisted to disk or the database. Scoped to this
# server process, and every read is additionally scoped to the uploading
# person_id (see get_upload) so one user can never fetch another's upload_id.
_UPLOADS: dict[str, EphemeralUpload] = {}


def _guess_mime(ext: str) -> str:
    ext = ext.lower().lstrip(".")
    return {
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "png": "image/png",
        "gif": "image/gif",
        "webp": "image/webp",
        "bmp": "image/bmp",
    }.get(ext, "application/octet-stream")


def parse_upload(filename: str, content: bytes, person_id: str) -> EphemeralUpload:
    """Parse an uploaded file into ephemeral text chunks and images.

    Supports:
    - Plain images (png/jpg/gif/webp/bmp): stored as a single image for the
      vision model to analyze directly.
    - .pptx: split PER SLIDE into text frames (-> text_chunks, answered by
      the text drafting model with the same citation-verification loop as
      governed content) and picture shapes (-> images, answered by the
      vision model). This is the actual "classify text vs image, per slide"
      behavior -- python-pptx's shape tree already carries that distinction
      natively, so it's read directly rather than guessed at.
    """
    upload_id = uuid.uuid4().hex
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    upload = EphemeralUpload(upload_id=upload_id, filename=filename, person_id=person_id)

    if ext in IMAGE_EXTENSIONS:
        upload.images.append(
            EphemeralImage(
                label=filename,
                base64_data=base64.b64encode(content).decode("ascii"),
                mime_type=_guess_mime(ext),
            )
        )
    elif ext == "pptx":
        prs = Presentation(io.BytesIO(content))
        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_text_parts: list[str] = []
            image_count = 0
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                    slide_text_parts.append(shape.text_frame.text.strip())
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image_count += 1
                    image = shape.image
                    upload.images.append(
                        EphemeralImage(
                            label=f"Slide {slide_idx}, Image {image_count}",
                            base64_data=base64.b64encode(image.blob).decode("ascii"),
                            mime_type=_guess_mime(image.ext),
                        )
                    )
            if slide_text_parts:
                upload.text_chunks.append(
                    EphemeralTextChunk(label=f"Slide {slide_idx}", text="\n".join(slide_text_parts))
                )
    else:
        raise ValueError(f"Unsupported file type '.{ext}'. Supported: images (png/jpg/gif/webp/bmp) and .pptx.")

    _UPLOADS[upload_id] = upload
    return upload


def get_upload(upload_id: str, person_id: str) -> Optional[EphemeralUpload]:
    """Retrieve a previously parsed upload, scoped to the uploading person_id."""
    upload = _UPLOADS.get(upload_id)
    if upload is None or upload.person_id != person_id:
        return None
    return upload


def discard_upload(upload_id: str, person_id: str) -> None:
    """Explicitly drop an upload from memory (e.g. user removes the attachment)."""
    upload = _UPLOADS.get(upload_id)
    if upload is not None and upload.person_id == person_id:
        _UPLOADS.pop(upload_id, None)
